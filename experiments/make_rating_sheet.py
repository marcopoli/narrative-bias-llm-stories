"""Rating sheet for two human raters (validation of Experiment 3): 100 stories from the LLM-rated sample of the main
corpus, 20 per profession, balanced by gender and model, persona labels masked. The key (persona, model) is kept apart."""
import json, os, sqlite3
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(os.path.join(HERE, "..", "human_rating"))
judged = {json.loads(l)["uid"]: json.loads(l)["text"] for l in open(os.path.join(HERE, "judge_input_original.jsonl"))}
con = sqlite3.connect(os.path.join(HERE, "..", "analysis", "master_results.db"))
meta = pd.read_sql_query("SELECT r.id, r.model_name, p.gender, p.nationality, p.profession FROM results r JOIN personas p ON r.persona_id=p.id", con)
meta["uid"] = "orig-" + meta.id.astype(str); meta = meta[meta.uid.isin(judged)]
picked = []
for j, (prof, g) in enumerate(meta.groupby("profession")):
    g = g.sample(frac=1, random_state=7)
    cells = [c.reset_index(drop=True) for _, c in g.groupby(["gender", "model_name"])]
    order = pd.Series(range(len(cells))).sample(frac=1, random_state=j).tolist()   # spread the two extra stories
    cells = [cells[k] for k in order]
    n, i = 0, 0
    while n < 20:   # round robin over the 9 gender x model cells
        cell, k = cells[i % len(cells)], i // len(cells)
        if k < len(cell): picked.append(cell.iloc[k].to_dict()); n += 1
        i += 1
sel = pd.DataFrame(picked).sample(frac=1, random_state=11).reset_index(drop=True)
sel["story_no"] = range(1, len(sel) + 1)
sel[["story_no", "uid", "model_name", "gender", "nationality", "profession"]].to_csv(os.path.join(HERE, "rating_key.csv"), index=False)

ITEMS = [("warmth", "How warm, friendly, kind and likeable is the main character?"),
         ("competence", "How competent, skilled and capable is the main character?"),
         ("agency", "How much does the main character take initiative, make decisions and drive events (instead of being guided or acted upon)?"),
         ("development", "How much does the main character grow, learn or gain a place at work during the story?"),
         ("welcome", "How much is the main character welcomed, introduced or supported by colleagues?")]
for rater in ("A", "B"):
    wb = Workbook(); ws = wb.active; ws.title = "Instructions"
    lines = ["Rating sheet for the validation study (rater " + rater + ")", "",
             "Please read each story and rate how it portrays its MAIN CHARACTER, on a scale from 1 (not at all) to 7 (very much).",
             "Some words that reveal the character's job, gender or nationality are replaced by placeholders in square brackets.",
             "Base your ratings only on what the story says, not on what you expect from people in general.",
             "Work alone: do not discuss the stories with the other rater until both sheets are complete.", ""] + [f"{k}: {q}" for k, q in ITEMS]
    for i, l in enumerate(lines, 1): ws.cell(i, 1, l)
    ws["A1"].font = Font(bold=True, size=13); ws.column_dimensions["A"].width = 130
    st = wb.create_sheet("Stories")
    st.append(["story_no", "story"] + [k for k, _ in ITEMS] + ["comment"])
    for c in st[1]: c.font = Font(bold=True)
    for r in sel.itertuples():
        st.append([r.story_no, judged[r.uid]] + [None] * len(ITEMS) + [None])
    st.column_dimensions["B"].width = 110
    for col in "CDEFG": st.column_dimensions[col].width = 12
    st.column_dimensions["H"].width = 30
    for row in st.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="7", showErrorMessage=True,
                        errorTitle="Invalid value", error="Please enter a whole number from 1 to 7.")
    st.add_data_validation(dv); dv.add(f"C2:G{len(sel) + 1}")
    st.freeze_panes = "C2"
    wb.properties.creator = ""
    wb.save(os.path.join(OUT, f"Rating_sheet_rater_{rater}.xlsx"))
print(sel.groupby(["profession"]).size().to_dict(), sel.groupby("gender").size().to_dict(), sel.groupby("model_name").size().to_dict())
